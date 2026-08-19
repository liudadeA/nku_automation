from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def main():
    path = sorted(Path(".").glob("*.docx"))[1]
    doc = Document(str(path))
    paras = doc.paragraphs

    p45 = paras[45]
    p47 = paras[47]
    p52 = paras[52]
    p53 = paras[53]
    p58 = paras[58]
    p61 = paras[61]
    p66 = paras[66]
    p70 = paras[70]
    p73 = paras[73]
    p77 = paras[77]
    p79 = paras[79]

    p45.text = (
        "系统围绕“感知-定位-导航-抓取-返航”构建完整任务链路：用户通过语音或键盘发出"
        "启动命令后，状态机首先触发定点导航；机器人到达目标区域后进入原地搜索，调用"
        "YOLOv8识别瓶子目标；检测稳定后切换到视觉伺服模块，对目标进行对中和距离调整；"
        "目标进入抓取位姿后，由Dynamixel五轴机械臂完成伸出、下探、闭合和抬起；任务"
        "结束后，系统再返回起点或设定终点。Ubuntu原生方案与WSL2跨平台方案共享这一"
        "总体流程。"
    )

    p47.text = (
        "对于WSL2实现，系统进一步采用“Windows感知端 + WSL2控制端”的分层架构。"
        "Windows主机负责Kinect RGB/深度采集、YOLOv8目标检测以及深度图到LaserScan"
        "的向量化转换，再通过自定义双端口TCP协议将检测结果和扫描数据发送给WSL2中"
        "的ROS节点；WSL2端则专注于状态机调度、AMCL定位、move_base导航、视觉伺服与"
        "机械臂控制。这样的设计既规避了部分驱动在WSL2中的兼容问题，也降低了控制侧"
        "的实时计算压力。示意图见图2.3。"
    )

    p52.text = (
        "系统采用中央状态机作为调度核心。state_machine.py 订阅用户命令、导航状态、"
        "目标检测结果、瓶子到位信号和机械臂到位信号，并发布 /control_state 与 "
        "/controll_msg 统一驱动各功能节点。这样做可以保证导航、搜索、视觉伺服、抓取"
        "和返航严格按照状态顺序执行，避免多个节点同时争用底盘控制权。速度指令经 "
        "cmd_vel_mux 仲裁：导航阶段主要使用 /cmd_vel_mux/input/navi，搜索、视觉伺服"
        "和返航阶段则切换到 /cmd_vel_mux/input/teleop。"
    )

    p53.text = (
        "在模块接口层面，感知数据主要通过 /detection_result、/bottle_caught 和 "
        "/scan 三类话题流动。AMCL利用 /scan 完成粒子滤波定位，navi.py 基于 move_base "
        "的 action 接口发送目标点，search_spin.py 在搜索阶段提供原地旋转，"
        "visual_approach.py 根据检测框中心与高度信息生成闭环控制指令，jicheng.py "
        "通过五个关节控制话题完成机械臂动作，back_2_zero.py 则在任务末端完成返航。"
        "Ubuntu方案直接使用ROS生态中的驱动和传感器转换链路；WSL2方案增加了一层TCP"
        "通信适配，但对上层ROS接口保持兼容，因此两种实现可以共享绝大部分控制逻辑。"
    )

    p58.text = (
        "在Ubuntu原生方案中，机器人结合Kinect深度数据与底盘里程计信息完成室内建图。"
        "系统先将深度图转换为 /scan 激光数据，再利用Gmapping执行二维栅格地图构建。"
        "Gmapping本质上采用Rao-Blackwellized粒子滤波同时估计机器人轨迹与环境地图，"
        "适合本项目这种结构化室内场景。通过保存建图结果，后续任务执行时可以直接加载"
        "地图，减少重复建图带来的时间开销。如下图3.1。"
    )

    p61.text = (
        "在导航阶段，两套方案都采用 AMCL + move_base 的经典ROS导航栈。启动后 "
        "global_localize.py 会向 /initialpose 发布大协方差位姿，使粒子在地图中充分"
        "散布；机器人原地旋转和短距离移动后，AMCL粒子逐步收敛到真实位姿。随后 "
        "navi.py 通过 actionlib 向 move_base 发送目标点，并记录出发位置供返航使用；"
        "move_base 结合全局路径规划与局部避障输出导航速度，实现从起点到目标区域的"
        "稳定移动。如下图3.2。"
    )

    p66.text = (
        "视觉抓取链路由目标检测、目标筛选、视觉伺服和机械臂控制四部分组成。首先，"
        "YOLOv8输出候选目标框后，bottle_filter.py 仅保留 bottle 类结果；随后 "
        "visual_approach.py 对检测框中心横向偏差和目标框高度进行闭环调节。控制过程"
        "分为旋转对准、前后粗调、前后微调、二次旋转对准和短距离开环前进等阶段，并在"
        "目标稳定满足阈值后通过 /bottle_in_pos 触发下一步动作。如下图3.3。"
    )
    insert_paragraph_after(
        p66,
        (
            "机械臂部分采用5个Dynamixel舵机实现四段位姿抓取：位姿1为预抓取伸出，"
            "位姿2下探靠近目标，位姿3闭合夹爪，位姿4抬起搬运。各位姿之间通过插值方式"
            "平滑过渡，降低舵机突变带来的振动风险。抓取完成后，系统还会将视觉伺服阶段"
            "录制的底盘速度序列反向回放，使机器人先撤回到导航目标附近，再调用 "
            "move_base 返回起点，从而形成较完整的闭环作业流程。"
        ),
        style=p66.style,
    )

    p70.text = (
        "系统在多个层面实现了鲁棒性设计。感知层面，目标检测结果经过类别筛选和多帧"
        "稳定判断，原始版本使用“10帧中至少5帧命中”确认目标，WSL2版本则通过状态机与"
        "筛选器联合降低误检影响；视觉伺服节点还对检测框中心和高度引入滑动平均与速度"
        "平滑，避免底盘出现突兀抖动。导航层面，各节点仅在被状态机激活时才发布控制指令，"
        "能够有效防止多源速度冲突；搜索阶段通过原地旋转扩大观测范围，返航阶段则结合"
        "录制回放和 move_base 二次导航提高回收成功率。通信层面，WSL2方案的双端口TCP"
        "协议加入握手、序号、MD5校验与断线重连机制，启动脚本还会提前检测服务器连通性，"
        "从而提升跨平台系统在真实运行中的稳定性与可维护性。"
    )

    sec34 = insert_paragraph_after(p70, "3.4 系统的质量保障方案", style=paras[69].style)
    insert_paragraph_after(
        sec34,
        (
            "系统的质量保障主要从模块划分、参数管理和分层验证三个方面展开。首先，导航、"
            "搜索、视觉伺服、机械臂、通信和语音交互都被封装为独立ROS节点，并以固定话题"
            "和状态机接口连接，便于单模块替换、联调和问题定位。其次，目标点、代价地图"
            "参数、舵机等待时间、服务器地址等关键配置均通过 launch 文件或 yaml 文件集中"
            "管理，减少硬编码带来的维护风险。最后，项目提供 full_pipeline.launch、"
            "full_debug.launch、manual_gmapping.launch 以及仿真启动文件等多种运行入口，"
            "支持从单模块验证到整链路回归测试的渐进式开发流程，体现了机器人软件工程中的"
            "可维护性、可测试性和可复用性要求。"
        ),
        style=p70.style,
    )

    sec4_body1 = insert_paragraph_after(
        p73,
        (
            "在本次项目中，本人主要完成了大作业文件夹下基于 Ubuntu 原生环境的核心功能"
            "开发与系统集成工作。具体包括搭建 full_pipeline.launch 所对应的整套运行链路，"
            "将底盘驱动、Kinect 感知、depthimage_to_laserscan、AMCL 定位、move_base 导航、"
            "YOLOv8 检测、机械臂控制以及语音交互等模块组织成可直接联调的完整抓取系统。"
        ),
        style=p79.style,
    )
    insert_paragraph_after(
        sec4_body1,
        (
            "此外，我重点完成了 controll_core.py 中央状态机、visual_approach.py 视觉"
            "伺服靠近逻辑，以及抓取与返航流程的联调优化。其中状态机实现了瓶子检测防抖、"
            "阶段切换和模块协调；视觉伺服实现了基于检测框中心和高度的旋转对准、粗调、"
            "微调与开环前进；同时还参与了 bottle_filter.py、jicheng.py 和 back_2_zero.py "
            "等节点的整体联调与测试。通过这些工作，我对机器人软件工程中的模块划分、接口"
            "设计和系统级调试有了更深入的理解。"
        ),
        style=p79.style,
    )

    insert_paragraph_after(p77, "5.1 项目总结", style=paras[44].style)
    p79.text = (
        "本项目围绕“移动底盘导航 + 视觉感知 + 机械臂操作”的课程目标，完成了一套能够"
        "在真实室内环境中运行的 TurtleBot 瓶体自主抓取系统。项目不仅实现了 Ubuntu "
        "原生环境下的建图、定位、导航与抓取闭环，也探索了 WSL2 条件下的跨平台部署方案，"
        "通过将 Windows 端感知与 WSL2 端ROS控制解耦，降低了传感器驱动和推理计算对控制"
        "系统的耦合影响。从工程结果看，系统已经具备较清晰的模块划分、较稳定的通信机制"
        "和较完整的任务流程，体现了需求分析、模块设计、接口定义、联调验证和问题修复等"
        "机器人软件工程核心能力。"
    )

    doc.save(str(path))
    print(f"updated: {path.name}")


if __name__ == "__main__":
    main()
