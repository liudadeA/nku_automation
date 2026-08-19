#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def md_to_docx(md_file_path, docx_file_path):
    """
    将Markdown文件转换为DOCX文件
    支持基本的Markdown语法：标题、列表、加粗、水平线
    """
    # 创建Word文档
    doc = Document()
    
    # 设置全局字体
    for style in doc.styles:
        if style.name in ['Heading 1', 'Heading 2', 'Heading 3', 'Normal']:
            style.font.name = 'Microsoft YaHei'
            
    # 读取Markdown文件
    with open(md_file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    # 处理每一行
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # 跳过空行
        if not line:
            i += 1
            continue
        
        # 处理标题
        if line.startswith('#'):
            # 确定标题级别
            level = 0
            while line[level] == '#' and level < len(line):
                level += 1
            
            # 提取标题文本
            title_text = line[level:].strip()
            
            # 添加标题到文档
            if level == 1:
                doc.add_heading(title_text, level=1)
            elif level == 2:
                doc.add_heading(title_text, level=2)
            elif level == 3:
                doc.add_heading(title_text, level=3)
            
            i += 1
            continue
        
        # 处理水平线
        if line == '---':
            doc.add_paragraph().add_run('—' * 40).font.size = Pt(1)
            i += 1
            continue
        
        # 处理列表项
        if line.startswith('* '):
            # 开始一个新的列表
            items = []
            while i < len(lines) and (lines[i].strip().startswith('* ') or not lines[i].strip()):
                list_line = lines[i].strip()
                if list_line.startswith('* '):
                    items.append(list_line[2:])
                i += 1
            
            # 添加列表到文档
            if items:
                for item in items:
                    # 处理加粗文本
                    if '**' in item:
                        p = doc.add_paragraph()
                        parts = item.split('**')
                        for j, part in enumerate(parts):
                            if j % 2 == 0:
                                p.add_run(part)
                            else:
                                p.add_run(part).bold = True
                    else:
                        doc.add_paragraph(item, style='List Bullet')
            continue
        
        # 处理普通段落
        if line.startswith('**'):
            # 处理加粗标题
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            parts = line.split('**')
            for j, part in enumerate(parts):
                if j % 2 == 0:
                    p.add_run(part)
                else:
                    run = p.add_run(part)
                    run.bold = True
                    run.font.size = Pt(12)
        else:
            # 处理普通文本
            p = doc.add_paragraph(line)
        
        i += 1
    
    # 保存文档
    doc.save(docx_file_path)
    print(f"转换完成：{docx_file_path}")


if __name__ == "__main__":
    # 命令行参数解析
    parser = argparse.ArgumentParser(description='将Markdown文件转换为DOCX文件')
    parser.add_argument('input_file', help='输入的Markdown文件路径')
    parser.add_argument('output_file', help='输出的DOCX文件路径')
    args = parser.parse_args()
    
    # 调用转换函数
    md_to_docx(args.input_file, args.output_file)