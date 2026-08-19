#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
使用python-markdown和python-docx库将带有HTML样式的Markdown转换为DOCX
支持HTML内联样式
"""

import markdown
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from bs4 import BeautifulSoup


def html_to_docx(html_content, doc):
    """
    将HTML内容转换为DOCX文档
    
    Args:
        html_content: HTML内容
        doc: Document对象
    """
    soup = BeautifulSoup(html_content, 'html.parser')
    
    for element in soup.children:
        if element.name == 'h1':
            # 处理一级标题
            heading = doc.add_heading(level=0)
            heading_run = heading.add_run(element.text)
            
            # 应用HTML样式
            if element.get('style'):
                style = element.get('style')
                if 'font-family' in style:
                    font_family = style.split('font-family:')[1].split(';')[0].strip()
                    if '宋体' in font_family:
                        heading_run.font.name = 'Times New Roman'
                        heading_run.font._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                if 'font-size' in style:
                    font_size = style.split('font-size:')[1].split(';')[0].strip()
                    if font_size == '16pt':
                        heading_run.font.size = Pt(16)
                if 'font-weight' in style and 'bold' in style.split('font-weight:')[1].split(';')[0].strip():
                    heading_run.font.bold = True
            
            if element.get('align') == 'center':
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
        elif element.name == 'h2':
            # 处理二级标题
            heading = doc.add_heading(level=2)
            heading_run = heading.add_run(element.text)
            
            # 应用HTML样式
            if element.get('style'):
                style = element.get('style')
                if 'font-family' in style:
                    font_family = style.split('font-family:')[1].split(';')[0].strip()
                    if '黑体' in font_family:
                        heading_run.font.name = 'Times New Roman'
                        heading_run.font._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                if 'font-size' in style:
                    font_size = style.split('font-size:')[1].split(';')[0].strip()
                    if font_size == '14pt':
                        heading_run.font.size = Pt(14)
                    elif font_size == '13pt':
                        heading_run.font.size = Pt(13)
                if 'font-weight' in style and 'bold' in style.split('font-weight:')[1].split(';')[0].strip():
                    heading_run.font.bold = True
            
            if element.get('align') == 'center':
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
        elif element.name == 'h3':
            # 处理三级标题
            heading = doc.add_heading(level=3)
            heading_run = heading.add_run(element.text)
            
            # 应用HTML样式
            if element.get('style'):
                style = element.get('style')
                if 'font-family' in style:
                    font_family = style.split('font-family:')[1].split(';')[0].strip()
                    if '黑体' in font_family:
                        heading_run.font.name = 'Times New Roman'
                        heading_run.font._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                if 'font-size' in style:
                    font_size = style.split('font-size:')[1].split(';')[0].strip()
                    if font_size == '12pt':
                        heading_run.font.size = Pt(12)
                if 'font-weight' in style and 'bold' in style.split('font-weight:')[1].split(';')[0].strip():
                    heading_run.font.bold = True
                    
        elif element.name == 'p':
            # 处理段落
            para = doc.add_paragraph()
            para_run = para.add_run(element.text)
            
            # 应用HTML样式
            if element.get('style'):
                style = element.get('style')
                if 'font-family' in style:
                    font_family = style.split('font-family:')[1].split(';')[0].strip()
                    if '宋体' in font_family:
                        para_run.font.name = 'Times New Roman'
                        para_run.font._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                if 'font-size' in style:
                    font_size = style.split('font-size:')[1].split(';')[0].strip()
                    if font_size == '12pt':
                        para_run.font.size = Pt(12)
                    elif font_size == '10.5pt':
                        para_run.font.size = Pt(10.5)
                if 'line-height' in style:
                    line_height = style.split('line-height:')[1].split(';')[0].strip()
                    if line_height == '1.5':
                        para.paragraph_format.line_spacing = 1.5
            
            if element.get('align') == 'center':
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                

def md_to_docx(md_file_path, docx_file_path):
    """
    将带有HTML样式的Markdown文件转换为DOCX文件
    
    Args:
        md_file_path: Markdown文件路径
        docx_file_path: 输出DOCX文件路径
    """
    # 创建Document对象
    doc = Document()
    
    # 设置默认样式
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal_style.paragraph_format.line_spacing = 1.5
    
    # 读取Markdown文件内容
    with open(md_file_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # 将Markdown转换为HTML
    html_content = markdown.markdown(md_content, extensions=['extra'])
    
    # 将HTML转换为DOCX
    html_to_docx(html_content, doc)
    
    # 保存DOCX文件
    doc.save(docx_file_path)
    print(f"转换完成：{md_file_path} -> {docx_file_path}")


if __name__ == "__main__":
    # 转换3_with_format.md到3_with_format.docx
    md_file = "3_with_format.md"
    docx_file = "3_with_format.docx"
    md_to_docx(md_file, docx_file)
