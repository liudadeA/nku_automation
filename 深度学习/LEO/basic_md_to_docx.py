#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
基本的Markdown到DOCX转换器
"""

import sys
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def basic_md_to_docx(md_path, docx_path):
    """基本的Markdown到DOCX转换"""
    # 创建文档
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # 读取Markdown文件
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 按段落分割
    sections = content.split('\n\n')
    
    in_code_block = False
    code_content = []
    
    for section in sections:
        if not section:
            continue
        
        lines = section.split('\n')
        
        for line in lines:
            line = line.strip()
            
            if line.startswith('```'):
                # 代码块开始/结束
                if in_code_block:
                    # 代码块结束
                    if code_content:
                        # 添加代码段落
                        paragraph = doc.add_paragraph()
                        paragraph_run = paragraph.add_run('\n'.join(code_content))
                        paragraph_run.font.name = 'Consolas'
                        paragraph_run.font.size = Pt(10)
                        in_code_block = False
                        code_content = []
                else:
                    # 代码块开始
                    in_code_block = True
                continue
            
            if in_code_block:
                # 收集代码内容
                code_content.append(line)
                continue
            
            if line.startswith('# '):
                # 标题1
                heading = doc.add_heading(level=1)
                heading_run = heading.add_run(line[2:])
                heading_run.font.size = Pt(24)
                heading_run.bold = True
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith('## '):
                # 标题2
                heading = doc.add_heading(level=2)
                heading_run = heading.add_run(line[3:])
                heading_run.font.size = Pt(18)
                heading_run.bold = True
            elif line.startswith('### '):
                # 标题3
                heading = doc.add_heading(level=3)
                heading_run = heading.add_run(line[4:])
                heading_run.font.size = Pt(14)
                heading_run.bold = True
            elif line.startswith('#### '):
                # 标题4
                heading = doc.add_heading(level=4)
                heading_run = heading.add_run(line[5:])
                heading_run.font.size = Pt(12)
                heading_run.bold = True
            elif line.startswith('- ') or line.startswith('* '):
                # 无序列表
                paragraph = doc.add_paragraph(style='List Bullet')
                paragraph_run = paragraph.add_run(line[2:])
            elif line.startswith('1. '):
                # 有序列表
                paragraph = doc.add_paragraph(style='List Number')
                paragraph_run = paragraph.add_run(line[3:])
            elif line.startswith('!['):
                # 图片
                img_desc_end = line.find('](')
                img_path_start = img_desc_end + 2
                img_path_end = line.find(')', img_path_start)
                if img_desc_end != -1 and img_path_end != -1:
                    img_path = line[img_path_start:img_path_end]
                    if os.path.exists(img_path):
                        doc.add_picture(img_path, width=Inches(6.0))
            else:
                # 普通段落
                paragraph = doc.add_paragraph()
                paragraph_run = paragraph.add_run(line)
        
        # 在段落之间添加空行
        if not in_code_block:
            doc.add_paragraph()
    
    # 保存文档
    doc.save(docx_path)
    print(f"成功将 {md_path} 转换为 {docx_path}")


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("用法: python basic_md_to_docx.py input.md output.docx")
        sys.exit(1)
    
    input_path = sys.argv[1]
    output_path = sys.argv[2]
    
    if not os.path.exists(input_path):
        print(f"错误: 输入文件 {input_path} 不存在")
        sys.exit(1)
    
    basic_md_to_docx(input_path, output_path)
