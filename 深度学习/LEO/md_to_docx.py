#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将Markdown文件转换为DOCX格式
"""

import sys
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import markdown2
from bs4 import BeautifulSoup


def md_to_html(md_content):
    """将Markdown转换为HTML"""
    html_content = markdown2.markdown(md_content, extras=[
        'fenced-code-blocks',
        'tables',
        'footnotes',
        'header-ids',
        'code-friendly',
        'toc',
        'strike'
    ])
    return html_content


def html_to_docx(html_content, docx_path):
    """将HTML转换为DOCX"""
    # 创建文档
    doc = Document()
    
    # 解析HTML
    soup = BeautifulSoup(html_content, 'html.parser')
    
    # 遍历HTML内容
    for element in soup.children:
        if element.name == 'h1':
            # 添加标题1
            heading = doc.add_heading(level=1)
            heading_run = heading.add_run(element.text)
            heading_run.font.size = Pt(24)
            heading_run.bold = True
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif element.name == 'h2':
            # 添加标题2
            heading = doc.add_heading(level=2)
            heading_run = heading.add_run(element.text)
            heading_run.font.size = Pt(18)
            heading_run.bold = True
        elif element.name == 'h3':
            # 添加标题3
            heading = doc.add_heading(level=3)
            heading_run = heading.add_run(element.text)
            heading_run.font.size = Pt(14)
            heading_run.bold = True
        elif element.name == 'p':
            # 添加段落
            paragraph = doc.add_paragraph()
            paragraph_run = paragraph.add_run(element.text)
            paragraph_run.font.size = Pt(12)
        elif element.name == 'ul' or element.name == 'ol':
            # 添加列表
            for li in element.find_all('li'):
                paragraph = doc.add_paragraph(style='List Bullet' if element.name == 'ul' else 'List Number')
                paragraph_run = paragraph.add_run(li.text)
                paragraph_run.font.size = Pt(12)
        elif element.name == 'table':
            # 添加表格
            rows = element.find_all('tr')
            if rows:
                # 创建表格
                table = doc.add_table(rows=0, cols=len(rows[0].find_all(['th', 'td'])), style='Table Grid')
                
                # 处理表头
                header_cells = rows[0].find_all(['th', 'td'])
                header_row = table.add_row()
                for i, cell in enumerate(header_cells):
                    header_row.cells[i].text = cell.text
                    header_row.cells[i].paragraphs[0].runs[0].bold = True
                
                # 处理表格内容
                for row in rows[1:]:
                    cells = row.find_all(['th', 'td'])
                    table_row = table.add_row()
                    for i, cell in enumerate(cells):
                        table_row.cells[i].text = cell.text
        elif element.name == 'pre':
            # 添加代码块
            code_block = element.find('code')
            if code_block:
                # 添加代码段落
                paragraph = doc.add_paragraph(style='Normal')
                paragraph_run = paragraph.add_run(code_block.text)
                paragraph_run.font.name = 'Consolas'
                paragraph_run.font.size = Pt(10)
        elif element.name == 'img':
            # 添加图片（仅支持本地图片）
            img_src = element.get('src')
            if img_src and os.path.exists(img_src):
                doc.add_picture(img_src, width=Inches(6.0))
        elif element.name == 'blockquote':
            # 添加引用
            paragraph = doc.add_paragraph(style='Intense Quote')
            paragraph_run = paragraph.add_run(element.text)
            paragraph_run.font.size = Pt(12)
    
    # 保存文档
    doc.save(docx_path)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print("用法: python md_to_docx.py input.md output.docx")
        sys.exit(1)
    
    input_path = sys.argv[1]
    output_path = sys.argv[2]
    
    if not os.path.exists(input_path):
        print(f"错误: 输入文件 {input_path} 不存在")
        sys.exit(1)
    
    # 读取Markdown文件
    with open(input_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # 转换为HTML
    html_content = md_to_html(md_content)
    
    # 转换为DOCX
    html_to_docx(html_content, output_path)
    
    print(f"成功将 {input_path} 转换为 {output_path}")
